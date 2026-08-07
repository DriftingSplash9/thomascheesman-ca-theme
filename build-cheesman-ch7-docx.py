# -*- coding: utf-8 -*-
"""Build Cheesman-Ch7-Cooking-Years.docx — the cooking years unfolded from
Thomas's raw notes (Your-Words-Needed annotations) into his voice, for his
edit pass before it lands on the Cheesman long-read Chapter Seven.
Grounded to his notes only; the items he marked private are excluded."""
import docx
from docx.shared import Pt, RGBColor

d = docx.Document()
style = d.styles['Normal']
style.font.name = 'Georgia'
style.font.size = Pt(12)

GREY = RGBColor(0x88, 0x88, 0x88)
RED = RGBColor(0x9b, 0x2c, 0x2c)

def title(t): d.add_heading(t, level=0)
def chap(t): d.add_heading(t, level=1)
def para(t):
    p = d.add_paragraph(t)
    p.paragraph_format.space_after = Pt(10)
    return p
def note(t):
    p = d.add_paragraph()
    r = p.add_run(t); r.italic = True; r.font.color.rgb = GREY; r.font.size = Pt(10)
    return p
def flag(t):
    p = d.add_paragraph()
    r = p.add_run('FLAG — ' + t); r.italic = True; r.font.color.rgb = RED; r.font.size = Pt(10)
    return p

title('The cooking years — Cheesman long-read, Chapter Seven')
note('Unfolded from your notes into the page voice — nothing invented, two things '
     'deliberately left out (see the flags). When you send it back I’ll lay it onto '
     'the Cheesman long-read’s Chapter Seven, which currently covers all of this in '
     'three sentences. Once the prose is approved, I’ll pull your 29 new WP uploads '
     'into a describe-sheet so you can caption them, and weave the best through the '
     'chapter.')
flag('The Ric’s Grill paragraph refers to the previous chef’s departure only as '
     '“a devastating loss in his own family.” Your notes were specific, but this '
     'is a public page about identifiable people in a small city — my advice is to '
     'leave the specifics out, and the sentence grieves with them without naming it. '
     'Your call.')
flag('The bank/corporate item you marked (keep private) is excluded entirely. The '
     'shrimp profanity is tamed but the joke is kept. A few beats here (Township 71’s '
     'close, the college class) also appear in brief on your own page — I’ve angled '
     'them differently so the two pages rhyme rather than repeat; read both before '
     'approving.')

d.add_paragraph()

chap('The cooking years')

para('My cooking life started in February of 2001, in the last semester of my second '
     'year of college. My dad had fallen ill and money for school was tight, and I knew '
     'a few people who worked at The Keg and spoke well of it — and I knew the place '
     'myself, if I’m honest, from the Friday-night martinis: three ounces for $3.99, '
     'which is exactly the kind of mathematics a college budget understands. I applied '
     'to serve, because my restaurant experience to that point had all been front of '
     'house, and they hired me instead as the salad tender on a station called Hub — '
     'bread, salads, and desserts, parked where you could see and talk to everyone in '
     'the building. If a thing was worth doing it was worth doing right; that was the '
     'whole of my philosophy, and Hub was where I started doing it.')

para('Then one night not a single appetizer cook came to work. I’d been watching '
     'those clowns long enough to have the station mostly figured out, so the general '
     'manager and I jumped on it together and rocked a busy night harder than any '
     'proper pair of cooks had in a long while — and it wasn’t long before I was '
     'pulled off Hub and parked on appies for good. So much for good deeds. Within a '
     'year I felt trapped there: I wanted to learn the whole kitchen, and managers can '
     'be slow to recognize talent and desire, so I begged and pleaded my way into '
     'training on the other stations one at a time. I had a proper blowout with my chef '
     'in the middle of it, too — he wanted me in on my day off when I was already '
     'sick, and I let him have it, because I was not going to be taken advantage of. It '
     'turned out to be pneumonia, and I worked straight through it while carrying my '
     'courses. Cranky, sick, tired, and done with it — but I came out the other side '
     'trained, and the kitchen began to feel like a thing I could grow into.')

para('Grande Prairie broke every rule corporate ever wrote. In the cities you sell a '
     'dessert to one table in ten and an appetizer to one in five; up here four tables '
     'in five took an appetizer, every second guest drank, and every second guest '
     'ordered some kind of seafood, because the town was full of east-coasters making '
     'killer money with no family around to spend it on. Corporate’s whole model was '
     'built on stagnation, and we would not stop growing — record sales and record '
     'counts, year over year — so we bent what needed bending. The shrimp is my '
     'favourite example. We sold so much of it that peeling raw shrimp would have eaten '
     'a cook’s entire shift — eight hours of one precious human life spent picking '
     'at shrimp, and nobody alive should spend their hours that way — so we quietly '
     'bought them peeled and deveined, ran the stock down whenever corporate came to '
     'visit, and took a special delivery of contraband shrimp once they’d gone. We '
     'did the same with the calamari, which eventually went out through the proper '
     'channels to every Keg in the country. You’re welcome.')

para('My chef was getting older, and one day he didn’t come to work; it turned out '
     'his heart had tried to quit on him. Somebody had to steer the ship, and I '
     'didn’t want him stressing about it from a hospital bed, so I put it in high '
     'gear and took on the whole job, because nobody else knew how or what to do — '
     'and to my surprise they gave me all the rope I asked for. I’d been made '
     'assistant kitchen manager about six months before, when the man holding the job '
     'moved on to run another kitchen. By the time chef was well enough to come back, '
     'his kitchen was a new place: busier than ever, some new faces, and the jobs '
     'themselves rewritten — the prep above all, because it doesn’t matter what '
     'you do on the line at night if the food made for you that day was junk; there is '
     'no rescuing it. We grew the whole time I was there, all the way to line-ups out '
     'the door four nights a week, eighty people deep on a Friday.')

para('What ended it wasn’t the work; it was me. Between the kitchen and the '
     'Hajdu-Cheney my feet were wearing out and my shoulders were starting to go, and '
     'there was a girlfriend and a child on the way, and I knew I could not give that '
     'kitchen the attention it needed and still be the man my family was about to '
     'need. So I left it with people I trusted to run it their own way. About seven '
     'years on, the place closed for a big renovation under a new operating partner, '
     'and I lost touch with most of the old crew — absorbed, by then, into a new '
     'life and a new kitchen of my own.')

para('After a six-month stint cooking at a local pub, I went in as head chef at a '
     'place called Ric’s Grill — my first day, as it happens, was the day my '
     'daughter Patience was born — and I walked into a heartbreak. The chef before '
     'me had left after a devastating loss in his own family, the kind no family '
     'should ever have to carry, and his mom was still there, holding the place '
     'together because somebody had to. It took her a while to trust that I could '
     'carry it instead, and the transition was nothing like smooth, and I was patient '
     'with every bit of it, because what they had been through buys a person all the '
     'patience I’ve got. The franchise was coming apart underneath us at the same '
     'time: the fax line cut off and reclaimed in our own name on our own dime, a '
     'point-of-sale system we had to source ourselves, an owner far away telling me '
     'what oil to fry in rather than trusting a chef to run his kitchen — and as '
     'the other stores fell like dominoes, every gift card the chain had ever sold '
     'seemed to find its way to our till. We honoured them right up to the day we shut '
     'Ric’s down and reopened as something new.')

para('The something new was Township 71, and the timing could not have been crueller. '
     'Oil collapsed, the economy slid for the first time in local memory, and the '
     'design itself was a mismatch — a room built around downtown-office lunches, '
     'in a town that does not have downtown-office people. What Grande Prairie has is '
     'oilfield workers in Cat boots and coveralls, rolling in at six or seven, looking '
     'to fill up for a decent price. Oh, how I wish I knew then what I know now. In '
     'the middle of all this I was also teaching for the college — eight students '
     'in a culinary course written just for them, held right in the restaurant, so '
     'that my class watched one place close, gut itself, and open as a brand-new '
     'concept, with a weekly lab in my kitchen where we built appetizers and soups and '
     'entrées and once served a whole cocktail party while Ric’s was still '
     'open. Madness, some weeks. What a class. Township 71 lasted nine long, hard '
     'months, two of them profitable, and the owners made the call. The staff meeting '
     'was brutal — there were people in that room who had been with the owner far '
     'longer than I had — and I felt I had failed every one of them, because the '
     'decision was Emma’s to make but I had been a key part of how she came to it. '
     'I wish the HCS had let me give that place what I gave the Keg. But my young '
     'family needed me, and they were going to need me long after my working years '
     'were done.')

para('Daniel was born right after we closed Township 71, and I took a couple of '
     'months to be present for the early days. Before I left, a man named Rob came to '
     'talk me into working for him, and one tour later I understood why: he had leased '
     'the kitchen in a hotel and was running banquets, buffets, a bar, and a dining '
     'room out of it — a whole different animal from running line cooks, and '
     'exactly the new experience I wanted. That was Majors, and I worked it full-time '
     'until 2019, when I had to stop for good. I had exhausted the medical treatments '
     'for pain, and past a certain point a kitchen is a dangerous place to be '
     'inattentive in, and so is a highway. My feet were growing spikes out of their '
     'soles, my jaw was shortening, my shoulders were collapsing; shingles, a blood '
     'clot, a lump removed, a dental operation, all inside a few short years. I was '
     'falling apart on a schedule.')

para('I loved Majors, and its people hold a place in my heart for good. Rob had a '
     'gift for gathering the kindest, warmest people into one building, and even '
     'after I stopped I kept helping out through that first year — the money '
     'helped, and I think I was a help too. I had Daniel in preschool and Patience in '
     'grade one by then, two or three school runs a day, so it was nothing to swing by '
     'and do the ordering or whatever the kitchen needed. Most of the staff are '
     'strangers to me now, and Rob passed away from illness last winter. I will '
     'remember that great man for the rest of my life. And Jenn, from those years, '
     'holds an honour nobody else has managed: hers is the only house my kids have '
     'ever spent the night at — a badge she can wear proudly, because my kids are '
     'picky about their company.')

para('These days the cooking is family-sized. I tried a few in-home dinners and wine '
     'dinners afterward, but three or four hours on these feet and hands was more than '
     'they had to give. So now I cook for my own, on my own time, and the best of it '
     'is the teaching: Patience or Daniel getting curious about how a dish comes '
     'together, and Faith starting to hover at the counter — she helps make the '
     'tacos, sometimes. It is a smaller kitchen than I used to run. The team is '
     'better.')

out = r'C:\Users\thoma\Desktop\Heritage research\Cheesman-Ch7-Cooking-Years.docx'
d.save(out)
print('saved:', out)
