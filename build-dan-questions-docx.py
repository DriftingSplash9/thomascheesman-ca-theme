# -*- coding: utf-8 -*-
"""Build Dan-Interview-Questions.docx — the question list Thomas asked for,
to record a sit-down with Papa Dan (Dan Haiste) while he's still here.
Covers Jim Haiste, Sydney, the Rycroft in-laws, the Steinkes, and Lana.
Modeled on the 20-question list Thomas sent his own father."""
import docx
from docx.shared import Pt, RGBColor

d = docx.Document()
style = d.styles['Normal']
style.font.name = 'Georgia'
style.font.size = Pt(12)

GREY = RGBColor(0x88, 0x88, 0x88)

def title(t): d.add_heading(t, level=0)
def chap(t): d.add_heading(t, level=1)
def note(t):
    p = d.add_paragraph()
    r = p.add_run(t); r.italic = True; r.font.color.rgb = GREY; r.font.size = Pt(10)
    return p
def q(t):
    p = d.add_paragraph(t, style='List Number')
    p.paragraph_format.space_after = Pt(6)
    return p

title('Questions for Papa Dan')
note('For a sit-down recording — a phone voice memo on the table works fine. '
     'Don’t worry about order or completeness; stories beat facts, and dates only '
     'if they come easily. One good wandering answer is worth ten short ones. '
     'Anything he gives feeds four pages: the Haiste long-read (Jim’s chapter is '
     'the acknowledged placeholder), the Rycroft long-read (Eric has exactly two '
     'facts), the Steinke page (the thinnest line), and the family pages.')

chap('About your dad — Jim Haiste')
q('What did your dad actually do for work — the trades, the employers, the jobs he talked about? Walk me through it year by year if you can.')
q('Did he have a phrase or a saying everyone remembers him by?')
q('What did he do with his hands when he wasn’t working — a shop, a garden, cards, a rink?')
q('What’s the one story about your dad that gets told at every family gathering?')
q('How did he carry the Parkinson’s years — anything he kept doing right to the end?')
q('How did he and your mom meet?')
q('What kind of grandfather was he to Melanie?')

chap('About Sydney — your grandfather')
q('What was Sydney like as a father, from what your dad said?')
q('Did anyone talk about the move up from Saskatchewan — the dust years, what they left, why the Peace?')
q('Where did Sydney die, and where is he buried? (Our research turned up a Sydney Haiste buried at Assiniboia, Saskatchewan, d. 1983 — we had him living out his days in Alberta, so one of those is wrong and you may be the only person who knows which.)')
q('Did Jim or Sydney ever mention the grandfather who came from England — Ernest? Or his wife — does the name Alice ring any bell? (The records point to an Alice Maud Smith of Leeds, married to Ernest in 1903.)')

chap('About the Rycroft in-laws')
q('What was Sam Rycroft like as a father-in-law?')
q('Was the “born in Honolulu” fact ever told at the family table — Sam’s grandfather Eric, born under a Hawaiian sky? How did Sam tell it?')
q('What do you remember about Eric Rycroft yourself — anything at all: how he farmed, what he drove, how he spoke?')
q('What was Bette like in her prime — before the last years?')
q('Did Bette ever talk about her parents, Henry and Martha Steinke — their farm at Sexsmith, their church, the language at home, where in Germany the family came from?')
q('Do you know which of Bette’s siblings (or their kids) are still around to ask?')

chap('About Lana — only if you feel like it')
q('How did you and Lana meet?')
q('What was she proudest of?')
q('What should Melanie’s kids know about their grandma Lana that only you can tell them?')

note('Last one, for the end of the tape: “What question should I have asked you '
     'that I didn’t?” — it’s the one that always finds the best story.')

out = r'C:\Users\thoma\Desktop\Heritage research\Dan-Interview-Questions.docx'
d.save(out)
print('saved:', out)
