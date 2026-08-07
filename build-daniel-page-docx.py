# -*- coding: utf-8 -*-
"""Build Daniel-Page-Text.docx — the full /family/daniel prose in Thomas's
voice (de-listed, grounded to his answers, nothing invented) for him to edit.
Headings = chapter titles; [Photo: ...] italic lines = caption notes."""
import docx
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

d = docx.Document()
style = d.styles['Normal']
style.font.name = 'Georgia'
style.font.size = Pt(12)

def title(t):
    h = d.add_heading(t, level=0)
def chap(t):
    d.add_heading(t, level=1)
def para(t):
    p = d.add_paragraph(t)
    p.paragraph_format.space_after = Pt(10)
    return p
def note(t):
    p = d.add_paragraph()
    r = p.add_run(t); r.italic = True; r.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
    r.font.size = Pt(10)
    return p
def cap(t):
    note('[Photo: ' + t + ']')

title('Daniel — /family/daniel page text')
note('Edit this freely — it’s the prose for Daniel’s page, de-listed and put into your '
     'voice, grounded to your answers with nothing invented. The bold lines are the chapter '
     'headings (they also become the left-hand table of contents on the page). The grey '
     '[Photo: …] lines are the caption under each picture — tweak those too if you like. '
     'When you’re done, send it back and I’ll lay it onto the live page.')
note('Eyebrow: Son — middle   ·   Subtitle: Charlie Brown — the quiet observer, the maker, '
     'the boy with a constellation on his face')

d.add_paragraph()

# ---------- The arrival (opening, no heading on the page) ----------
chap('The arrival')
cap('the elegant son — hero')
para('Daniel was born on the thirtieth of June, 2015, at the QE2 in Grande Prairie — about '
     'average on the scales, a little short on the tape, and he’s carried that compact build '
     'all the way to today. Patience had already made us a family; Daniel rounded it out, a boy '
     'and a girl, and just like that we were complete — the envy, I think, of all the families '
     'who’d ended up with only boys, or only girls, or no children at all. I would have sworn my '
     'heart was full after the first one. Daniel showed me there was a whole other room in it I '
     'hadn’t found yet.')
para('By the time he came along, Melanie and I were as ready as two people get. I happened to be '
     'between jobs, so I took a couple of months off after the birth and got to be the kind of '
     'help I never quite managed to be when Patience arrived. He turned up just after Township 71 '
     'closed its doors — the whole pregnancy ran almost exactly the length of that place’s last '
     'stretch, which is the sort of thing you only notice years later. Because Patience had come '
     'by C-section, his was booked the same way; once you’ve had the one, they tell me, a natural '
     'birth gets risky enough that they just schedule the next. And the day we brought him home '
     'there was a nice surprise waiting: both sets of grandparents, and my own dad, who’d driven '
     'up without telling us.')
para('My first memory of him is holding him right after his first weigh-in. Melanie was still '
     'woozy and getting stitched up from the operation, so the very first snuggle fell to me — and '
     'he was so small, and so pink, with these enormous brown puppy eyes I’ve never quite gotten '
     'over. Patience would hold him too, and I swear it was the cutest thing I’d ever seen, every '
     'single time.')
cap('held-by-big-sister')
para('The first weeks were the hard kind nobody prints on a card. A tongue-tie meant '
     'breastfeeding just wouldn’t take, so instead there was pumping — endless pumping, a big '
     'machine we rented from the pharmacy that roared away in the night, cups and hoses and '
     'Melanie running on willpower alone, trying everything natural and medicinal she could find '
     'to keep the milk coming. When she should have been sleeping she was tending pumps. I was '
     'deep into my own pain by then and short on rest, and between the two of us we were about as '
     'worn out as people get. She had it worse, though. She always did, in those early days.')
para('We named him Daniel for a few reasons all at once. It was the name of my best friend '
     'growing up; it was the name of Melanie’s best friend from her childhood; and it was her '
     'father’s name on top of that. For some reason it had always been parked in the back of my '
     'mind as the name I’d use if I ever had a son — it was even on the shortlist for Patience, '
     'right up until we learned she was a she. His full name is Daniel Eric Thomas Cheesman: Eric '
     'for his great-grandfather, and Thomas for reasons I’ll keep to myself.')
para('He never did crawl, not properly. He butt-scootched instead — sat bolt upright and hauled '
     'himself across the floor one leg at a time, a bit like a rook sliding across a chessboard, '
     'and he could really move when there was a cup of milk waiting at the far end of it. Milk was '
     'his fuel right up until about six: first thing every morning, in cups and sippy cups, never '
     'a bottle. I have never in my life met such a milk fiend.')
cap('VIDEO — the butt-scootch in action')
para('He turned up bald — gloriously, completely bald, with a head sitting up in the ninetieth '
     'percentile — and his uncle took one look and christened him Charlie Brown. The name stuck '
     'around for a while. The head, he eventually grew into.')

# ---------- The quiet observer ----------
chap('The quiet observer')
cap('daniel-little')
para('Ask me to sum Daniel up and the words that come are compassionate, creative, and patient — '
     'and right behind them, shy, quiet, and private. He’s not the one who stands out in a crowd '
     'or works the room; he’s the boy off to one side, quietly taking the whole thing in and '
     'working out how the world really runs underneath whatever everyone’s saying about it. It '
     'showed up early, back in pre-K, when making friends and speaking up didn’t come easily to '
     'him. He takes his time coming out of his shell — but when he does, and he does, there’s a '
     'funny, sharp, adventurous kid in there, fully switched on.')
para('He’s got three small moles under his right eye, set in a line, and they are Orion’s Belt to '
     'the life — even the brightness of them matches the stars. I noticed it years ago and I’ve '
     'never been able to unsee it: a boy with a constellation on his face, mad about space. I don’t '
     'think the two things are unrelated.')
para('His humour is dry and a little odd. He’ll let out a strange noise from nowhere, for no '
     'reason I’ve ever been able to pin down. But he’s at his funniest when he’s excited — when a '
     'story takes hold of him and he has to get it out right now, all in a rush, which is usually '
     'at bedtime, of course, the moment he’s finally ready to talk. He answers to a small pile of '
     'names around here — G, and Bro, and Brosky, and Bra, and Dude, and Little Dude, depending on '
     'the day. And when he’s tired and gone slow and whiny, I’ll catch myself calling it draggin’ '
     'ass — a phrase I’m sure he could do without.')

# ---------- The maker ----------
chap('The maker')
para('The thing you most need to know about Daniel is that he makes things. He’s a real artist — '
     'drawings, cartoons, comics, little stop-motion Lego animations he builds and films himself. '
     'Hand him a free afternoon with no rules attached and he’ll draw, or build something, or '
     'vanish into a game, and honestly I’m happy with any of the three.')
cap('terry-fator — Dawson Creek, to see the great Terry Fator, a real working puppeteer')
para('But it’s the puppets that have his whole heart right now. He’s gone deep on them, and on '
     'Fugglers — those gleefully ugly things with the human teeth — and he sews his own Kermits '
     'from scratch, cutting and stitching and gluing until there’s a whole cast of them lying '
     'about. Ask him what he wants to be when he grows up and the answer comes back with no '
     'hesitation at all: a puppeteer. I love that. It’s not a thing most ten-year-olds would even '
     'think to want, which is exactly why it’s so completely him.')
para('It started long before Kermit, too. Years ago he was already making little home movies with '
     'an Elmo, running the whole show himself — and if you watch a few seconds of one you can see '
     'the puppeteer he’s turning into, already in there.')
cap('VIDEO — one of his Elmo productions, before Kermit came along')
para('This spring he spent two or three weekends building what he calls the Man Hut — a fort '
     'thrown together out of old pallets and a tired girly playhouse we had going spare. The next '
     'weekend we painted it. Then the weather turned cool and rainy, which the land badly needed '
     'and the boy badly didn’t, and the Man Hut’s been waiting out the wet ever since. It’ll get '
     'its summer.')
cap('reading-together — I miss reading to you, son')
para('His reading habits are their own small comedy: Garfield, and then, with a completely '
     'straight face, the dictionary or the thesaurus. He likes words the way he likes Lego — as '
     'pieces you can pull apart and put back together a different way.')

# ---------- Bacon, metal, and a perfect day ----------
chap('Bacon, metal, and a perfect day')
cap('bubble-party — a bubble party at the neighbours’, thanks Graham and Cate')
para('His idea of a perfect day starts outside with his best friend. The two of them have secret '
     'knocks worked out, so when the right rhythm lands on the door Daniel lights up and bolts to '
     'answer it. There’d be a fire in that perfect day, and s’mores, and hotdogs, and very likely '
     'a bubble party somewhere in the mix. He’s an outdoor kid for the warm two-thirds of the '
     'year and a sensible indoor one for the rest of it, this being Alberta.')
para('If you ever need to get on Daniel’s good side, lead with bacon. It is, as best I can tell, '
     'the great love of his eating life — the thing he’d happily take at breakfast, lunch and '
     'supper if the world allowed it, and his answer to just about any kind of bad day. Feeling '
     'rotten? Bacon. It’s practically what the doctor orders in this house. After that he runs to '
     'the honest favourites of a boy his age — fried chicken, a pile of wings, a steak when he can '
     'talk me into one — with one hard, unexplained line drawn straight through tacos, which he '
     'simply will not eat. His music’s been creeping heavier the older he gets; these days it’s '
     'metal in his headphones — Slipknot, and a good deal of Eminem — though he’ll still throw on '
     'some Ozzy or Michael Jackson, and not so long ago he was a five-year-old playing Imagine '
     'Dragons on repeat. His colour, if you ask, is green and has been forever; his animal is a '
     'dog and only ever a dog. And his favourite film is Real Steel — the one about the boxing '
     'robots — which made me grin the first time I noticed, because for a winter there boxing was '
     'his thing too, and he was good at it.')
cap('PAIR — two of his great loves: Lucas the dog, and Allister the lizard')
para('The best place he’s ever been is British Columbia — the mountains and the lakes and the '
     'valleys of it. And we do it his way, which is the slow way: pull over at every lake, stop '
     'for every view, take the long road on purpose.')
cap('VIDEO — rocking out to AC/DC (Thunderstruck, if this page had a sound)')

# ---------- Brave in the ways that count ----------
chap('Brave in the ways that count')
para('Daniel feels things deeply, and that he gets straight from his mom — the pair of them '
     'welling up at the sad part of the same film while Patience and I quietly trade looks across '
     'the room. He says I love you easily, he’s a hugger, and when something’s got him low the '
     'cure is simple and it’s a snuggle. None of that is softness in the weak sense. I’ve watched '
     'him prove it.')
para('When I had my neck fused — a big, frightening operation, the kind that shuts a lot of kids '
     'right down — Daniel stayed steady. He didn’t panic. He asked his questions, he handed out '
     'hugs, he stayed close by. He has this uncommon knack for being near pain without backing '
     'away from it, and he carries my bad days more gently than I tend to carry them myself.')
para('He’s grown up around people whose minds run fast and switch tracks without much warning, '
     'and he meets them with patience and grace every time — rolls with it, changes the game on a '
     'dime, never once makes anyone feel like they’re a problem to be managed. That’s a rare thing '
     'in a grown adult. In a boy his age it stops me in my tracks.')

# ---------- Brother, middle, peacemaker ----------
chap('Brother, middle, peacemaker')
cap('helping-with-chores — already pitching in around the yard')
para('Daniel’s our middle child and our only boy, tucked in between two sisters about twenty-one '
     'months out on either side — which, as it happens, is almost the exact spacing my own '
     'brothers and I grew up with. He’s the helper of the house, and the peacemaker. He does the '
     'dishes, he pitches in without being asked, and when the temperature climbs between his '
     'sisters he’s usually the one who quietly brings it back down.')
para('He and Patience are close — Roblox together, fishing together, a game of catch in the yard — '
     'and about the only thing she does that truly gets under his skin is disappearing into her '
     'room for hours at a stretch. I’ve also caught her quietly coaching him on his hair and his '
     'clothes, sorting him out before school, and watching that happen does something to me: if '
     'they’re looking after each other like this now, I have to believe they’ll be alright when '
     'they’re grown and it really counts.')
cap('daniel-and-poppy — Daniel with Princess Poppy')
para('With Faith he’s a good big brother — reluctantly sometimes, the way big brothers are, but '
     'he’ll fold her into whatever he and his buddy are building, and he’s fierce about '
     'protecting her. The best of it happens when I’m not looking: the three of them making a '
     'meal together in the kitchen, sure I can’t hear a thing, while I sit just out of sight and '
     'let the whole scene play itself out. Those are the afternoons I’d keep, if I were only '
     'allowed to keep a few.')
cap('selfie-with-mom — a selfie with Mom')
para('He calls me Daaaad, with all the extra As, and he calls his mom Mom. He and his mom have '
     'their own things — a run to Subway, a crime drama side by side on the couch — and he and I '
     'have ours, which is Lego, and the little “blue jobs” around the house, and always, always '
     'the puppets. He’s close with his Papa Dan, who he flat-out adores, and with his Gramzie. '
     'And Christmas is his whole season: the food, the time off school, and that troublesome Elf '
     'who shows up every night for a month — who, I’m reliably told, is online now as well, and '
     'who I’d dearly love to leave me alone.')
para('For one of his birthdays we booked a place called the Sand Zone and invited his entire '
     'class — and very nearly all of them turned up. It turns out the quiet one is quietly '
     'popular. It didn’t surprise me in the least. People can feel his heart from clear across a '
     'room, even when he hasn’t said a word.')

# ---------- Who he's becoming ----------
chap('Who he’s becoming')
cap('off-to-school-2025 — off to school, 2025')
para('He’s in grade five at Crystal Park School, and he likes it there. Math is his subject — '
     'he’s proud of his division, and rightly so — and every year he brings home something or '
     'other for teamwork or for STEM. His report cards do the same quiet thing he does: they just '
     'keep getting a little better, year on year, no fuss made about it. He leans toward science '
     'and space, with a soft spot for a bit of history.')
para('Right now he’s between things. He boxed for a winter and was good at it, then set it down, '
     'and he’s taking his time working out what comes next — which strikes me as exactly the '
     'right way to be at ten. I don’t want him feeling he has to be anything in particular. I '
     'want him to find the thing he actually loves and chase it. He’s growing into a smart, kind, '
     'adventurous kid, and that soft, curious centre of him — the part that wants to know how the '
     'world really works — is the part I’d guard with everything I’ve got.')

# ---------- A letter to Daniel ----------
chap('A letter to Daniel')
para('Dear Daniel,')
para('I’m sure you already know that I love you — but I don’t think you’ll truly know it, not all '
     'the way down, until you’ve got a family of your own. So let me set it down here, where it’ll '
     'keep.')
para('I am proud of you. I’m proud of the way you handle your sisters and the ordinary weather '
     'of being a kid — the friendships, the school, the home, a houseful of pets — none of which '
     'is half as easy as grown-ups like to pretend they remember. You manage it like water off a '
     'duck’s back. Hold onto that. And hold onto the warm, caring heart you’ve got, because it’s '
     'rarer out there than it ought to be, and the world doesn’t always know what to do with the '
     'gentle ones. It can be cold, and it can be careless, and every so often it’ll drop '
     'something in your way that isn’t fair. Guard that softness anyway. I trust you to tell '
     'right from wrong, and to see clean through the nonsense when it turns up dressed as '
     'something important.')
para('You’ll be somewhere around thirty when you read this, and the honest truth is I’ve no idea '
     'what the world looks like from there, or what I’d tell you about it. All I can give you is '
     'what I see now. I think by then you’ll have found the love of your life — and I’d bet she’s '
     'a lot like you: full of wonder, curious about everything, up for the adventure.')
para('You make it all look so easy, buddy. Keep going, keep growing into your confidence, and '
     'never lose the part of you that makes things — puppets, comics, forts out of old pallets, '
     'whatever it’ll be by then.')
para('I love you. I always will.')
para('Dad')

cap('VIDEOS at the foot of the page (Home videos): the boxing-instructors beat-down, boxing with '
    '“Tooth Pick,” another Elmo production, and the ten-dollar marketplace train')

out = r'C:\Users\thoma\Desktop\Heritage research\Daniel-Page-Text.docx'
d.save(out)
print('saved:', out)
print('paragraphs:', len(d.paragraphs))
