"""Build Daniel-Get-To-Know.xlsx — a 100-question questionnaire for Thomas to
fill out as source material for Daniel's page. Navy/cream styling to match the
existing per-kid xlsx sheets. No formulas (it's a Q&A sheet)."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# --- palette (matches the kid sheets' navy/cream) ---
NAVY   = "1F3B5C"
NAVY2  = "2E4F73"
GOLD   = "E8C766"
CREAM  = "FBF3E0"
WHITE  = "FFFFFF"
INK    = "20303F"
FONT   = "Arial"

# (question, type)  — type is one of: "T / F", "Y / N", "Short answer"
TF, YN, SA = "T / F", "Y / N", "Short answer"

sections = [
    ("Section A · Getting to know Daniel", [
        ("What's Daniel's full name, and is there a story behind it (who he's named after, what it means)?", SA),
        ("When and where was Daniel born?", SA),
        ("What do you call Daniel day to day — any nicknames?", SA),
        ("In three words, how would you describe Daniel's personality?", SA),
        ("Daniel is more introvert than extrovert.", TF),
        ("Daniel makes friends easily.", TF),
        ("Daniel has a strong sense of humour.", TF),
        ("What makes Daniel laugh the hardest?", SA),
        ("Daniel is competitive.", TF),
        ("Daniel would rather be outdoors than indoors.", TF),
        ("What's Daniel's idea of a perfect day?", SA),
        ("Does Daniel love video games?", YN),
        ("What's Daniel's favourite game (video or otherwise) right now?", SA),
        ("Does Daniel play a sport?", YN),
        ("Which sports does Daniel play or love to watch?", SA),
        ("Does Daniel play a musical instrument?", YN),
        ("What music does Daniel like?", SA),
        ("What's Daniel's favourite food?", SA),
        ("Daniel is an adventurous eater.", TF),
        ("What food will Daniel absolutely refuse to eat?", SA),
        ("What's Daniel's favourite animal?", SA),
        ("Daniel is more of a dog person than a cat person.", TF),
        ("What's Daniel's favourite colour?", SA),
        ("What's Daniel's favourite movie or show?", SA),
        ("Does Daniel like to read? What does he reach for?", SA),
        ("Daniel is an early riser.", TF),
        ("Daniel keeps his room and things tidy.", TF),
        ("What's a quirk or habit that's unmistakably Daniel?", SA),
        ("What is Daniel genuinely great at?", SA),
        ("What does Daniel find hard or frustrating?", SA),
        ("Daniel is a patient kid.", TF),
        ("Daniel is shy around new people.", TF),
        ("What does Daniel say he wants to be when he grows up?", SA),
        ("Daniel loves building or making things with his hands.", TF),
        ("What's the best place Daniel has ever been?", SA),
        ("Does Daniel enjoy road trips and travelling?", YN),
        ("What's Daniel obsessed with or collecting right now?", SA),
        ("Daniel prefers summer to winter.", TF),
        ("What's Daniel's favourite holiday, and why?", SA),
        ("What's something Daniel is afraid of?", SA),
        ("What comforts Daniel when he's upset?", SA),
        ("With a free afternoon and no rules, what would Daniel do?", SA),
        ("What's a story about Daniel the family loves to retell?", SA),
        ("What's something most people don't realize about Daniel until they know him?", SA),
    ]),
    ("Section B · Daniel and his immediate family", [
        ("Where does Daniel fall in the birth order?", SA),
        ("Daniel is the only boy among the siblings.", TF),
        ("Daniel is close with his sister Patience.", TF),
        ("What do Daniel and Patience love to do together?", SA),
        ("Daniel and Patience butt heads often.", TF),
        ("What's something Patience does that drives Daniel up the wall?", SA),
        ("Daniel is protective of his sister Faith.", TF),
        ("What's Daniel like as a brother to Faith?", SA),
        ("Daniel and Faith play well together.", TF),
        ("Describe a sweet moment you've seen between Daniel and his sisters.", SA),
        ("What does Daniel call you?", SA),
        ("Daniel comes to you (Dad) first when something's wrong.", TF),
        ("What do you and Daniel love to do, just the two of you?", SA),
        ("What's the one lesson you most want Daniel to learn from you?", SA),
        ("What's something Daniel does that reminds you of yourself?", SA),
        ("What's something about Daniel that's entirely his own?", SA),
        ("Do you and Daniel share an inside joke? What is it?", SA),
        ("What's your proudest moment as Daniel's dad?", SA),
        ("What's the hardest part of being Daniel's dad?", SA),
        ("What does Daniel call his mom?", SA),
        ("What do Daniel and his mom love to do together?", SA),
        ("Daniel takes after his mom in some ways.", TF),
        ("In what ways is Daniel like his mom?", SA),
        ("How does Daniel show that he loves someone?", SA),
        ("Daniel says “I love you” easily.", TF),
        ("Daniel is affectionate — a hugger.", TF),
        ("What role does Daniel play in the family (peacemaker, clown, helper, leader)?", SA),
        ("Daniel pitches in around the house without being asked.", TF),
        ("What's a job or chore that's Daniel's responsibility?", SA),
        ("Is there a tradition or ritual Daniel shares with his siblings?", YN),
        ("What family trip or outing has Daniel loved the most?", SA),
        ("Daniel is the loudest one at the dinner table.", TF),
        ("What family thing does Daniel most look forward to?", SA),
        ("Is there extended family (grandparents, aunts, uncles) Daniel is especially close to?", YN),
        ("Who beyond the household does Daniel adore?", SA),
        ("How does Daniel handle big changes in the family?", SA),
        ("Daniel feels things deeply.", TF),
        ("What does Daniel need most from his family right now?", SA),
        ("What habit or value do you hope Daniel carries from this family into his own life?", SA),
        ("What's something about your family you hope Daniel always remembers?", SA),
    ]),
    ("Section C · Anything else for his page", [
        ("What grade is Daniel in, and where does he go to school?", SA),
        ("Daniel enjoys school.", TF),
        ("What's Daniel's favourite subject?", SA),
        ("What's a milestone or achievement Daniel is proud of?", SA),
        ("Has Daniel earned any awards, medals, or recognition?", YN),
        ("What's a goal Daniel is chasing right now?", SA),
        ("What's your earliest memory of Daniel, or his earliest of you?", SA),
        ("What do you hope Daniel's future holds?", SA),
        ("What's the one thing you'd want visitors to his page to know about him?", SA),
        ("What photo or moment captures Daniel perfectly?", SA),
        ("Is there a video of Daniel you'd want featured on his page?", YN),
        ("What's a catchphrase or thing Daniel always says?", SA),
        ("If Daniel's page had a signature colour, what would it be?", SA),
        ("If a song played on Daniel's page, what would it be?", SA),
        ("What would you want to say to Daniel in an open letter on his page?", SA),
        ("Anything else about Daniel the world should know?", SA),
    ]),
]

wb = Workbook()
ws = wb.active
ws.title = "Daniel — 100 Questions"
ws.sheet_view.showGridLines = False

thin = Side(style="thin", color="D9C9A3")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

# column widths
ws.column_dimensions["A"].width = 5
ws.column_dimensions["B"].width = 70
ws.column_dimensions["C"].width = 14
ws.column_dimensions["D"].width = 52

# --- title ---
ws.merge_cells("A1:D1")
t = ws["A1"]
t.value = "Getting to Know Daniel — 100 Questions"
t.font = Font(name=FONT, size=16, bold=True, color=WHITE)
t.fill = PatternFill("solid", fgColor=NAVY)
t.alignment = Alignment(horizontal="left", vertical="center", indent=1)
ws.row_dimensions[1].height = 34

# --- subtitle / note ---
ws.merge_cells("A2:D2")
s = ws["A2"]
s.value = ("Fill in the Your Answer column — you decide what's public, so skip "
           "anything too private. 44 + 40 + 16 = 100.")
s.font = Font(name=FONT, size=10, italic=True, color=INK)
s.fill = PatternFill("solid", fgColor=GOLD)
s.alignment = Alignment(horizontal="left", vertical="center", indent=1)
ws.row_dimensions[2].height = 22

# --- column headers ---
hdr = ["#", "Question", "Type", "Your Answer"]
for col, label in enumerate(hdr, start=1):
    c = ws.cell(row=3, column=col, value=label)
    c.font = Font(name=FONT, size=11, bold=True, color=WHITE)
    c.fill = PatternFill("solid", fgColor=NAVY)
    c.alignment = Alignment(horizontal=("left" if col == 2 else "center"),
                            vertical="center", indent=(1 if col == 2 else 0))
    c.border = border
ws.row_dimensions[3].height = 22
ws.freeze_panes = "A4"

row = 4
qnum = 0
zebra = True
for sec_title, questions in sections:
    # section banner
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    b = ws.cell(row=row, column=1, value=f"{sec_title}  ({len(questions)} questions)")
    b.font = Font(name=FONT, size=12, bold=True, color=GOLD)
    b.fill = PatternFill("solid", fgColor=NAVY2)
    b.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[row].height = 26
    for col in range(1, 5):
        ws.cell(row=row, column=col).border = border
    row += 1
    zebra = True

    for q, qtype in questions:
        qnum += 1
        fill = CREAM if zebra else WHITE
        zebra = not zebra

        num = ws.cell(row=row, column=1, value=qnum)
        num.font = Font(name=FONT, size=10, bold=True, color=INK)
        num.alignment = Alignment(horizontal="center", vertical="center")

        qc = ws.cell(row=row, column=2, value=q)
        qc.font = Font(name=FONT, size=10, color=INK)
        qc.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)

        tc = ws.cell(row=row, column=3, value=qtype)
        tc.font = Font(name=FONT, size=9, bold=(qtype != SA), color=(NAVY if qtype != SA else "8A7A52"))
        tc.alignment = Alignment(horizontal="center", vertical="center")

        ac = ws.cell(row=row, column=4, value="")
        ac.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)

        for col in range(1, 5):
            cell = ws.cell(row=row, column=col)
            cell.fill = PatternFill("solid", fgColor=fill)
            cell.border = border
        ws.row_dimensions[row].height = 30
        row += 1

out = r"C:\Users\thoma\Desktop\Heritage research\Daniel-Get-To-Know.xlsx"
wb.save(out)
print("saved:", out, "| total questions:", qnum)
