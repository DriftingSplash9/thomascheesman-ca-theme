# -*- coding: utf-8 -*-
"""Build Patience-Page-Text.docx — the CURRENT live /family/patience prose,
extracted verbatim from page-patience.php for Thomas to edit (his voice pass).
Headings = chapter titles; [Photo: ...] italic lines = caption notes."""
import docx
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

d = docx.Document()
style = d.styles['Normal']
style.font.name = 'Georgia'
style.font.size = Pt(12)

def title(t):
    h = d.add_heading(t, level=0)
def chap(t):
    d.add_heading(t, level=1)
def para(t):
    p = d.add_paragraph(t)
    p.paragraph_format.space_after = Pt(10)
    return p
def para_seg(segs):
    """segs = list of (text, italic) tuples — for paragraphs with <em> spans."""
    p = d.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    for text, ital in segs:
        r = p.add_run(text)
        r.italic = ital
    return p
def note(t):
    p = d.add_paragraph()
    r = p.add_run(t); r.italic = True; r.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
    r.font.size = Pt(10)
    return p
def cap(t):
    note('[Photo: ' + t + ']')

title('Patience — /family/patience page text')
note('This is the prose exactly as it stands on the live page today — nothing re-voiced, '
     'nothing added. Edit it freely: the bold lines are the chapter headings (they also '
     'become the left-hand table of contents on the page), and the grey [Photo: …] lines '
     'are the caption under each picture — tweak those too if you like. When you’re done, '
     'send it back and I’ll lay it onto the live page.')
note('Eyebrow: Daughter — eldest   ·   Subtitle: The natural-born leader who tries to '
     'hide her dimples')

d.add_paragraph()

# ---------- Opening (no heading on the page) ----------
chap('Opening (no heading on the live page)')
cap('Caught mid-yawn. Something in those eyes looked infinite to me — like the whole '
    'universe was hers if she wanted it. There’s an 11×18 print of this one '
    'somewhere in the house.')
para('Patience is my first miracle, and I don’t reach for that word lightly. I never '
     'thought I’d have kids — part of me thought I shouldn’t; Hajdu-Cheney made '
     'the whole question feel like a gamble I had no business making. Melanie made it '
     'anyway. And the moment they set her in my arms, the floor dropped out of everything '
     'I thought I understood about love: there was so much more of it in the world than '
     'anyone had told me, and I knew all at once how every ancestor down the long line of '
     'us must have felt holding their own. Life is a fluke and a miracle in the same '
     'breath, and we are all of us unspeakably lucky to be here.')
cap('Holding her, flat-out disbelieving I was somebody’s father.')
para('We called her Patience — Melanie’s idea, and I loved it the second I heard it. '
     'Her middle name is Kristi-Ann: Kristi-Ann for a dear friend of Melanie’s who '
     'was lost to a motorcycle, and Ann for my mom, Maryann. A name with two people '
     'already living inside it before she’d taken a breath — and she’s been '
     'growing to fill all three ever since.')
para('She was due on the fourteenth of September, and then she simply… wasn’t. '
     'The days went by one at a time — grandparents arriving and leaving again, the doctor '
     'beginning to murmur about inducing. The twenty-third finally sent us to the '
     'hospital, and the hospital sent us straight back home to wait some more. We '
     'returned on the twenty-fourth, were admitted, and when she wouldn’t dilate and '
     'her heart rate started to slip, the call came for an emergency caesarean. The '
     'twenty-fourth, naturally, was also the morning I was due to start a new job as head '
     'chef at Ric’s Grill. I’d hauled the orientation binders all the way to the '
     'maternity ward, figuring I’d get a little reading done between visits. I never '
     'turned a page. Some priorities sort themselves out for you.')
cap('Asleep over Mom’s shoulder — the spit of the slumped angel in her tattoo.')
para('That first night I sat up by the big windows with her while Melanie slept, the two '
     'of us a couple of days past the last real rest either of us had had. We’d just '
     'come back from her bath, and she was wrapped tight the way only a pediatric nurse '
     'can manage — a neat little parcel of a person. I looked out at the dark and felt '
     'the whole thing land at once: I was a father, she was healthy, we were going to be '
     'alright. Breathe in. Breathe out. (The feeding took us a while to crack — a tongue '
     'tie, an upper lip tie, a couple of snips that didn’t take, and in the end she '
     'settled happily on formula. First lesson of fatherhood: the plan is whatever '
     'works.)')

# ---------- Determined, self-driven, sensitive ----------
chap('Determined, self-driven, sensitive')
cap('First smiles.')
para('If you made me sum her up in three words, those are the ones — and the sensitive '
     'one shows first. She is shy, properly shy, the kind that floods her cheeks pink and '
     'tugs the dimples out at the smallest provocation. Turn the radio up, say hello to '
     'one of her friends, call her name across a parking lot, and there they are, those '
     'dimples, betraying her in front of everyone. She would much prefer they stayed '
     'hidden — which is, of course, the entire reason I can never quite stop trying to '
     'coax them out.')
cap('Where’s my coffee?')
para('She’s twelve going on sixteen now, which mostly means I see the back of a '
     'closing door. She lives online — gaming, chatting, watching shows with friends who '
     'are three rooms and one screen away — and I have taken to texting my own daughter '
     'to come and eat in my own kitchen. Weekends belong to sleepovers at her best '
     'friend’s — the same one since kindergarten, which is about as far back as a '
     'friendship goes — the all-night kind nobody will admit are all-nighters. And every '
     'single morning when I drop her off, she reaches over and turns the radio down so '
     'her friends won’t hear what her dad listens to. I crank it back up. I’m '
     'never fast enough; the door’s already shut. What she actually listens to once '
     'that door closes is a mystery to me — whatever’s climbing TikTok, I assume, '
     'though she’s not telling.')

# ---------- The one who runs the show ----------
chap('The one who runs the show')
cap('First day of kindergarten, 2018.')
para_seg([
    ('For all that vanishing-teenager act, hand her a job and she is suddenly, '
     'magnificently in charge. At 7:55 every school morning, on the dot, Patience takes '
     'command of the house — chasing the other two to find their things and make the '
     'truck, and not above shoving her dawdling old man along with them. Most mornings I '
     'just lean into it and let her run the operation; there’s no arguing with '
     'weather, and anyway I’d rather watch her become whoever it is she’s '
     'becoming. It’s the same out in the world: me in the wheelchair, cruising the '
     'mall, Patience herding her brother and sister around me — ', False),
    ('stay in line, quit dawdling, stop touching everything', True),
    (' — every word a parent has ever said, handed down a generation early. She’ll '
     'be a pro by the time she has her own.', False),
])

# ---------- The weight she carries ----------
chap('The weight she carries')
cap('Cuddles with Mary, late in her life. The most carefree dog — a truly kind soul who '
    'never minded the kids.')
para('But there’s a cost to being the one who holds it all together, and she’s '
     'the one who pays it. For all the take-charge, she is the most sensitive of my '
     'three, and she carries more than a girl her age ever should. Some of that is simply '
     'the lot of the eldest. Some of it, if I’m honest, is me. When I had my '
     'sledding accident, and then the spinal fusion, the worst of it came after the '
     'surgery — when I couldn’t talk, couldn’t do much of anything at all — and '
     'Patience was a rock. She got me through the hospital stay and a long stretch of '
     'what came after: a little girl doing a grown woman’s worrying. It still shows '
     'in small ways — on a Saturday she’ll quietly work down a list of chores so her '
     'mum doesn’t spend her one day off cleaning. Nobody asks her to.')
cap('Meeting her baby brother. She’s been a great big sister to Daniel ever since.')
para('I think she lost a little of her childhood early to that, and I couldn’t stop '
     'the caution that grew up in its place. To this day, if she hears a loud bang '
     'somewhere in the house, or a stumble on the stairs, her very first fear is that '
     'it’s me. If I could lift one thing off her shoulders, it would be exactly '
     'that. I wish she’d be gentler with herself and let the parents do the '
     'parenting — but the hardest trick in this whole life is learning when to set a '
     'thing down, and I’m trying to teach it to her the only way I know how, which '
     'is by example. Mixed results so far, I’ll admit.')

# ---------- The little things ----------
chap('The little things')
cap('Daddy’s helper. They don’t make chef whites much smaller than that.')
para_seg([
    ('She is, it has to be said, a creature of small and very particular appetites. Ten '
     'o’clock at night is when she comes alive in the kitchen — ice cream above all '
     '(I find the empty tubs in her wake like breadcrumbs), a muffin, a cookie, and when '
     'the cupboards let her down, raw cookie dough eaten with no shame whatsoever. Taco '
     'Pizza Pops and corn dogs hold the daytime fort. She’ll watch Sam and Colby '
     'like it’s coursework, keeps a quiet soft spot for Young Sheldon, and once '
     'handed her whole heart to Stranger Things. And she is careful with a dollar in a '
     'way I have never once managed to be — ', False),
    ('cheeping out', True),
    (', she calls it — careful enough that I’ve half a mind to sign over the family '
     'books and take an early retirement.', False),
])
cap('Buried to the chin after an afternoon of sandcastles.')
para('The two of us have always been good in a kitchen, if you grade on mess. There was '
     'the legendary snickerdoodle afternoon when I mixed up pounds and cups and wound up '
     'doubling a batch that was already doubled — we ate snickerdoodles for a fortnight, '
     'gave them away by the bagful, froze a stash, and mailed a box clear across the '
     'country to my mom and dad. Kitchen nil; Thomas and Patience, one. And long before '
     'any of that, in her very first spring, she pulled herself up to standing right '
     'about the time I committed the single dumbest act available to a new father: I '
     'pinched a blade of grass between my thumbs and blew, just to see what would happen, '
     'and she screamed like the world was ending. I felt like an ass for a week. I still '
     'do it now and then, God help me — except now she just rolls her eyes and goes '
     'pink.')

# ---------- On the road ----------
chap('On the road')
cap('Road tripping.')
para('Some of the moments I turn over most are the smallest ones — all three kids’ '
     'hair flying as we punched through the highway tunnels, shrieking with joy in the '
     'back seat for no reason except the dark and the speed of it. We’ve logged a '
     'lot of road, this family, and most of it blurs together kindly. One stretch '
     'won’t, though: a trip out to friends in Chase, B.C., where they parked us in a '
     'camper and our buddy Dan, restless to get the boat wet, took us out for a slow '
     'cruise one afternoon. I’ve been on every kind of vessel the Shuswap floats — '
     'canoe, houseboat, jet ski, the lot — and that aimless, sun-warmed hour beat every '
     'one of them, for no reason I can name except the company.')
cap('GP Stompede.')
cap('One day we’ll catch a fish. Right?')
para('Not every outing earns a postcard. There’s the Stampede most summers, and the '
     'fishing trips that have yet, against all available odds, to produce a single fish. '
     'And there was the long loop down to B.C. and back through Calgary when she was '
     'tiny — that one to lay her in her great-grandmother’s arms. I hadn’t '
     'watched Gramma Sandy hold a baby since my sister Amber was small, and there '
     'aren’t many photographs I’d run back into a fire for the way I would '
     'that one.')
cap('With her great-grandma, Sandy Cheesman.')
cap('Patience versus the toboggan, in three acts. Spoiler: the toboggan wins.')

# ---------- The big sister ----------
chap('The big sister')
cap('A rare, tender moment with Faith — kidding; they have more of these than either of '
    'them would ever admit, and I’ve got the proof.')
para('She is, above all, a big sister — and she wears it differently with each of them. '
     'With Daniel it comes easy: she keeps him pointed in the right direction, and the '
     'two of them trade the kind of ordinary brother-sister needling I treasure precisely '
     'because it’s so ordinary. With Faith it’s stormier, and I won’t '
     'pretend otherwise — Patience is stubborn as the day is long, and admitting '
     'she’s wrong ranks, in her estimation, somewhere just shy of dying. But every '
     'so often I round a corner and find them doing makeup together, or playing, or '
     'Patience quietly rallying everyone into a game online — and it knocks me sideways, '
     'in the best possible way. I didn’t make my peace with my own brothers until '
     'our twenties; we’re thick as thieves now. So I take the long view. One day, I '
     'have to believe, the girls will understand each other the way only sisters can.')
cap('The lemonade-and-iced-tea stand — a childhood must on a hot day.')

# ---------- What she's mastered ----------
chap('What she’s mastered')
cap('The Award of Excellence — twice. Once in grade four, and again in grade six.')
para('Then there’s the matter of the report cards. Top marks across the board, year '
     'after year, to the point where I honestly couldn’t tell you how she’d '
     'improve — though her teacher swears she’s only just beginning to blossom, '
     'which is a thrilling and faintly terrifying thing to hear about a kid already '
     'running the table. The Award of Excellence twice over: grade four, and again in '
     'grade six. And the funniest part is she has no idea what to do with a compliment — '
     'tell her she’s done something extraordinary and she’ll go pink and study '
     'her shoes. Modest right down to the bone, which is, of course, its own kind of '
     'extraordinary.')

# ---------- Now & next ----------
chap('Now & next')
cap('Besties.')
para('When I let myself picture her grown, the thing I hope above all is still burning is '
     'her kindness. She feels everything, and she’s gotten very good at not showing '
     'it — that tough exterior — but it lands hardest at home, where it can colour a '
     'whole day grey. School and friends and the sheer weight of being twelve are no '
     'small things to carry. So my one real worry is the plain one: that she’ll keep '
     'packing the load heavier than she needs to. And my comfort is that very same '
     'sentence read from the other end — a girl this determined, this self-driven, this '
     'kind, is going to be just fine. She always has been. I just get the privilege of '
     'watching.')

# ---------- A letter for later ----------
chap('A letter for later')
para('Hi Patience. So you’re twenty-five now, eh? More than anything, I hope you and '
     'your mum and your brother and sister are close — that’s the thing I wish '
     'hardest for you, because family is the one set of people worth keeping near your '
     'whole life long, and it’s worth the work.')
para('I’d put money on you being at a university, or maybe deep into a trade — with '
     'that head for numbers and the way you’ve always had things three steps figured '
     'out, I can see you in business or finance or economics, or somewhere none of us '
     'would have guessed. Whatever it is, I’m certain you chose it because you love '
     'it; you were never the type to settle for less. Be careful out there, mind — '
     'there’s a lot of garbage to wade through these days, more than there was in '
     'mine, and there was plenty then.')
para('Mostly I just want to know what these years were like. Do you and your mum still '
     'get into it over the silly things? Did you stay out chasing whatever’s new, or '
     'turn out the quieter one, more taken with history and how the world really works? '
     'Tell me everything. And know this, from the version of me writing it down at the '
     'kitchen table while you’re upstairs behind that door: I am, and I always have '
     'been, prouder of you than I have the words for. — Dad')

note('[After the prose, the live page continues with: the family cross-links, the '
     '139-photo year-sectioned photo wall, the 12-photo “More from the album” '
     'block, and the Home-videos section (Patience at the wedding 2016 · Patience '
     'and Daniel cuddling the kittens). Those aren’t in this document — just the '
     'prose and captions above.]')

out = r'C:\Users\thoma\Desktop\Heritage research\Patience-Page-Text.docx'
d.save(out)
print('saved:', out)
print('paragraphs:', len(d.paragraphs))
