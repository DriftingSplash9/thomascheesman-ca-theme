# -*- coding: utf-8 -*-
"""Generate fillable Q&A Word docs to prompt Thomas's writing for the kids'
pages + Melanie. Mirrors the Thomas.docx workflow: he types answers, I weave
them into the page prose."""
import os
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = r"C:\Users\thoma\Desktop\Heritage research"
os.makedirs(OUT, exist_ok=True)

INTRO = ("Answer as much or as little as you like, in your own voice — bullet points "
         "or full sentences both work. I'll weave your answers into the page prose "
         "(the same way we built your own “Thomasito” page). Skip anything you'd "
         "rather keep private; you decide what goes public. When a question brings a "
         "specific photo to mind, jot the filename from the photo-stories spreadsheet "
         "next to it.")

def build(title, subtitle, intro, sections, fname, privacy=None):
    doc = Document()
    # base style
    st = doc.styles["Normal"].font
    st.name = "Georgia"; st.size = Pt(11)

    h = doc.add_heading(title, level=0)
    sub = doc.add_paragraph(subtitle)
    sub.runs[0].italic = True; sub.runs[0].font.size = Pt(12)

    if privacy:
        p = doc.add_paragraph()
        r = p.add_run("A note on how this page will read:  ")
        r.bold = True; r.font.color.rgb = RGBColor(0x8a, 0x1c, 0x1c)
        r2 = p.add_run(privacy); r2.font.color.rgb = RGBColor(0x8a, 0x1c, 0x1c)

    ip = doc.add_paragraph(intro); ip.runs[0].font.size = Pt(10.5)
    doc.add_paragraph("")

    for sec_title, questions in sections:
        doc.add_heading(sec_title, level=1)
        for q in questions:
            qp = doc.add_paragraph()
            qr = qp.add_run(q); qr.bold = True
            # answer space
            ap = doc.add_paragraph()
            ar = ap.add_run("→ ")
            ar.font.color.rgb = RGBColor(0xb0, 0xb0, 0xb0)
            doc.add_paragraph("")
    out = os.path.join(OUT, fname)
    doc.save(out)
    print("SAVED", out)

# ---------- shared kid template ----------
def kid_sections(name):
    return [
        ("1 · The arrival", [
            f"What's the story of the day {name} was born? (where, when, what was going on in your life)",
            f"How did you choose the name {name}? Does it carry any meaning for you?",
            "What do you remember most about the first days and weeks?",
            f"What surprised you most about becoming {name}'s parent?",
        ]),
        ("2 · Who they are", [
            f"In three words, who is {name}?",
            "What's their personality like — and how has it shown itself since they were tiny?",
            "What makes them light up? What are they into right now?",
            "Something only someone who knows them well would notice?",
        ]),
        ("3 · The little things", [
            "A saying, mispronunciation, or phrase of theirs you never want to forget?",
            "A habit, ritual, or quirk that is so them?",
            "Favourite food / song / show / toy / place — then vs. now?",
            "What do they do that makes you laugh?",
        ]),
        ("4 · Moments that stuck", [
            "A milestone you'll always remember (first steps, first word, a first day...)?",
            "A small, ordinary moment you think about more than you expected to?",
            "A trip, holiday, or day out that stands out?",
            "A time they amazed you?",
        ]),
        ("5 · The brave stuff (optional)", [
            "A challenge they've faced, and how they handled it?",
            "A time they were braver or stronger than you expected?",
            "(Share only what you're comfortable making public — medical or sensitive details stay out unless you say otherwise.)",
        ]),
        ("6 · Them and us", [
            "How are they with their siblings? Any classic dynamics?",
            "Something the two of you share — an inside joke, an activity, a bond?",
            "A family tradition that's special with them?",
            "How are they like you? Like their mom? Like anyone back in the family lines?",
        ]),
        ("7 · Now & next", [
            "Who are they becoming?",
            "What are you most proud of?",
            "What do you hope for them?",
            f"If {name} reads this page in twenty years, what do you want them to know?",
        ]),
        ("8 · In their words / a note", [
            "Anything they've said about themselves you'd like to include?",
            f"A short note or letter to {name} you'd like printed on the page?",
        ]),
        ("9 · The photos", [
            f"Open {name.lower()}-photo-stories.xlsx — for any photo with a story, write what's happening in the 'Story' column and mark 'Y' to sprinkle it into the writing. Star your very favourites here too.",
        ]),
    ]

KIDS = {
    "Patience": ("Patience — writing prompts",
                 "The eldest. Your “first miracle” — the natural-born leader who tries to hide her dimples."),
    "Daniel":   ("Daniel — writing prompts",
                 "Your son."),
    "Faith":    ("Faith — writing prompts",
                 "The youngest — the one who, as you put it, came to you in a dream."),
}
for name, (title, subtitle) in KIDS.items():
    build(title, subtitle, INTRO, kid_sections(name), f"{name}-Questions.docx")

# ---------- Melanie ----------
mel_privacy = ("Melanie appears here as the mother of Patience, Daniel, and Faith and as part "
               "of the family's story — not as a current couple. Nothing will be framed "
               "romantically and nothing about your relationship status will be disclosed. You "
               "decide every word that goes public.")
mel_sections = [
    ("1 · Who she is", [
        "Where is Melanie from? A little about her background, her family, where she grew up.",
        "How would you describe her as a person — her personality, her strengths?",
        "What is she passionate about? What's she good at?",
        "Something people might not know about her?",
    ]),
    ("2 · As a mother", [
        "What kind of mother is she to the kids?",
        "A moment that captures her as a mom?",
        "What does each child get from her — traits, looks, temperament?",
        "You've said she took on the risk of children knowingly — how did she take to motherhood?",
    ]),
    ("3 · Moments & memories", [
        "A story or memory that really captures Melanie?",
        "A time she amazed you, or that you admired her?",
        "Something funny or warm that's so her?",
    ]),
    ("4 · The family you've built for the kids", [
        "What have the two of you built together for the children?",
        "A tradition, place, or routine that matters to the kids and to her?",
        "(Framed around the kids and the family — not the relationship.)",
    ]),
    ("5 · The photos", [
        "Do you have a set of photos of Melanie (with the kids or on her own) you'd like on the page? "
        "Send them over and I'll build her a gallery + a photo-stories spreadsheet like the kids'.",
    ]),
    ("6 · Anything else", [
        "What would you most like visitors to understand about Melanie?",
        "Anything you'd like to say to, or about, her on the page?",
    ]),
]
build("Melanie — writing prompts",
      "Mother of Patience, Daniel, and Faith.",
      INTRO, mel_sections, "Melanie-Questions.docx", privacy=mel_privacy)

print("DONE")
