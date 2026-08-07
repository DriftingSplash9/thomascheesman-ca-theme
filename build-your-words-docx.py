# -*- coding: utf-8 -*-
"""Build Your-Words-Needed.docx — the per-page punch list after the
2026-06-10 corpus editorial pass. Three kinds of items per page:
DONE (applied, FYI) / DECIDE (your call) / WRITE (only you can supply
the words — prompts for gaps and thin material)."""
import docx
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

d = docx.Document()
style = d.styles['Normal']
style.font.name = 'Georgia'
style.font.size = Pt(11)

GREY = RGBColor(0x88, 0x88, 0x88)
GREEN = RGBColor(0x3a, 0x7d, 0x44)
AMBER = RGBColor(0xb0, 0x8c, 0x4a)
RED = RGBColor(0x9b, 0x2c, 0x2c)

def chap(t):
    d.add_heading(t, level=1)

def note(t):
    p = d.add_paragraph()
    r = p.add_run(t); r.italic = True; r.font.color.rgb = GREY; r.font.size = Pt(10)
    return p

def item(tag, color, text):
    p = d.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(tag + '  '); r.bold = True; r.font.color.rgb = color
    p.add_run(text)
    return p

def done(t):  return item('DONE', GREEN, t)
def decide(t): return item('DECIDE', AMBER, t)
def write(t): return item('WRITE', RED, t)

d.add_heading('Your words needed — the per-page punch list', level=0)
note('Companion to Editorial-Review-2026-06-10.md. The mechanical corpus pass is '
     'applied and live (theme 1.0.502): em-dashes thinned, bodies de-bolded, research '
     'changelog moved to the Notes, shared endings de-duplicated, and the verified '
     'factual fixes made (the birthday gap is six days; Pieter’s baptism is 23 Aug 1783; '
     'the Rycroft line-roster now matches the Cheesman page; the McIver Notes markup is '
     'repaired). The corpus is roughly 1,500–2,000 words tighter with nothing lost but '
     'repetition. What remains is below — three kinds of item per page:')
note('DONE = applied, just so you know where to look on a read-through. '
     'DECIDE = a choice only you should make; tell Claude the verdict and it’s a '
     'one-line change. WRITE = a gap or thin spot only your memory can fill — '
     'send rough notes in any form (lists are fine!) and Claude will set them in '
     'your voice, grounded to what you give, nothing invented.')

# ---------------- Patience ----------------
chap('Patience — your DOCX pass (Patience-Page-Text.docx)')
note('These fold straight into the edit you’re doing now.')
decide('The Award of Excellence is told twice back-to-back — the photo caption and the '
       'prose paragraph both say “grade four, and again in grade six.” Cut it from the '
       'prose; let the caption carry the numbers.')
decide('“Pink” is the blush-marker three times (“floods her cheeks pink,” “rolls her '
       'eyes and goes pink,” “go pink and study her shoes”). Keep the first and last; '
       'vary the middle one.')
decide('“On the road” ends on two photo captions with no closing prose. Suggestion: move '
       'the toboggan triptych above the Gramma Sandy paragraph so “run back into a fire” '
       'closes the chapter — or write a one-line lander of your own.')
decide('mom vs mum — the page uses both (“my mom, Maryann” but “her mum doesn’t spend '
       'her one day off”). Pick one spelling, or deliberately split (mom = Maryann, one '
       'form for Melanie) — just make it a choice.')
decide('The show list (Sam and Colby / Young Sheldon / Stranger Things) still reads like '
       'the questionnaire. The ice-cream-at-ten and “cheeping out” material around it is '
       'fully alive — compress the list or fold one show into an anecdote.')
decide('The subtitle (“tries to hide her dimples”) front-runs the chapter-one dimples '
       'reveal. Optional tweak.')
write('The page is rich on baby-era and thinner in the now. One fresh twelve-year-old '
      'anecdote — a sleepover detail, a thing she said in the truck, the latest '
      '“cheeping out” move — would balance it. What happened recently that’s pure '
      'Patience?')

# ---------------- Daniel ----------------
chap('Daniel — your DOCX pass (Daniel-Page-Text.docx)')
note('One of these is a must-fix before the page ships; the rest fold into your edit.')
write('MUST-FIX: “I don’t think the two things are unrelated” (the moles / space line) '
      'is nearly word-for-word Faith’s live “I don’t think those two facts are '
      'unrelated” — and both are about space. The signature line can only sign one '
      'page; Faith’s is live, so rephrase Daniel’s in your words.')
decide('The boxing sentence appears twice (“for a winter there boxing was his thing too, '
       'and he was good at it” in Bacon-metal; “He boxed for a winter and was good at it” '
       'in Who he’s becoming). Cut the first — the second is doing real work about '
       'being “between things.”')
decide('The subtitle gives away both of the page’s best reveals (Charlie Brown AND the '
       'constellation). Trim it to one element.')
decide('“I have never in my life met such a milk fiend” mirrors Faith’s “I have never '
       'in my life met anyone with her engine.” Vary one of them.')
decide('The colour/animal favourites are pure questionnaire residue — cut them; give the '
       'Slipknot-to-Imagine-Dragons music arc its own sentence.')
decide('“Township 71 closed its doors” is unexplained on a standalone page — three or '
       'four words of identification, or save the coincidence for where Township 71 is '
       'actually told (your page).')
decide('“Just like that we were complete” collides with Faith’s “the last full stop on '
       'our family.” Time-scope it: “we felt complete.”')
decide('“I was deep into my own pain by then” — a first-time reader won’t know what it '
       'means for four more sections. Gloss it with a pointing word, or cut it (“She had '
       'it worse, though” lands harder without it).')
decide('The letter’s middle paragraph: “I’d bet she’s a lot like you” is the one '
       'greeting-card beat. Faith’s letter keeps it open (“Marry whoever makes you '
       'laugh”); consider matching that openness or cutting the partner bet.')
write('“Who he’s becoming” runs two paragraphs after six full chapters. One more '
      'concrete present-day Daniel — the thing he made this month, the current '
      'obsession, what he and his buddy are building right now — and the ending '
      'carries its weight.')

# ---------------- Faith ----------------
chap('Faith — live page (edits applied)')
done('Taco parenthetical cut; “and off the tablet” cut; the three words front-loaded '
     '(“Carefree, wandering, fireball…”); the duck paragraph split and untangled; the '
     'letter broken into three with “not a chef” as its own beat; the cat image promoted '
     'from caption to body; the Everest pile-up trimmed; the cookies lead-in now promises '
     'optimism, not velocity.')
write('OPTIONAL: hers is the shortest kid page (~1,950 words; her siblings run ~2,500). '
      'The missing beat is school — Patience and Daniel both have one. Send four or five '
      'facts (grade, what she loves, what she fights, a teacher line) and Claude will '
      'weave the chapter. Or leave her short and fierce — also a valid shape.')

# ---------------- Thomasito ----------------
chap('Thomasito — live page (edits applied)')
done('The wall confession now opens “Three kids I was told I’d never have” (it was '
     'splitting the neck-surgery chapter); the spark/treasure beat lands once, in the '
     'letter to the kids; the career “walls” are plain words now so the real wall owns '
     'the metaphor; the prophecy restatements trimmed to clauses; the cosmology paragraph '
     'cut to its anchor; the COVID aside and the “every strong person has a limit” '
     'dangler cut; the friends and concert roll-calls cut to the names that matter '
     '(Township 71’s list untouched); CH8’s newsletter sentences gone; em-dashes '
     '147 → 101. The page is ~430 words lighter, ~9,970 words.')
decide('The CH4 chapter title “Grande Prairie, and the wall” still says “wall” though '
       'its prose no longer uses the metaphor (headings were off-limits to the edit). '
       'Want it retitled? (e.g. something from the chapter’s own practicum/Keg/Tony '
       'material.)')
decide('The Mustang interlude could nest inside “The name I chose” (same era, same '
       'register) instead of standing alone — a bigger structural move, your call.')
decide('One line was cut as a pre-play of the letter: “one day… I’ll watch my own kids '
       'pass it on, and that will be the brightest flame of all.” Say the word if you '
       'want it back.')
write('CH6 — the kids and the wedding — is still the thinnest stretch for what it '
      'holds: Patience, Township 71, Daniel, the wedding, and Faith share ~700 words '
      'while the farms got 1,500. One small scene per kid (age 3–10, anything — a '
      'bath-time, a road moment, a thing they said) and one wedding detail would let '
      'the emotional centre breathe. Rough notes are fine.')
write('The concert list is now four acts. If one show MATTERED — a story, not a '
      'name — tell that one story and it earns its place back.')

# ---------------- Cheesman ----------------
chap('The Cheesmans — long-read (edits applied)')
done('The Mustang promise now points to your page (“I tell it over on my own page”, '
     'linked); Ch 7’s fourteen-years sentence is three sentences with Head-Chef-the-day-'
     'Patience-was-born standing alone; Ch 6 opens with an orienting clause and the '
     'Kool-Aid aside moved to Ch 2 beside Amber; Ch 4’s vignettes now escalate '
     '(friendly animals → labour → dangers, ending on Cone); the chosen/canal '
     'metaphor restated once instead of five times; eight “It was…” openers recast; '
     'the closing paragraph and restaurant names unbolded.')
decide('Keep the Kool-Aid aside in Ch 2, or cut it entirely? (One-line change.)')
decide('Your line kept its bold animal/farm/place names (Trumac, Candy Cane Farm…) — '
       'the other lines got a stricter people-only de-bold. Match them, or keep your '
       'page warmer? ')
write('Ch 7, the cooking years: fourteen years still go by in a blur. Two or three '
      'sentences each — in your words — on Ric’s Grill, teaching Culinary Arts, and '
      'the restaurant that didn’t make it (what was its name? what was the day it '
      'died like?) would turn a catalogue into chapters of a life.')

# ---------------- Docherty ----------------
chap('The Dochertys — long-read (edits applied)')
done('All “earlier draft” sentences out of the body (the Notes keep the forensics); '
     'the triple ending cut so “It came through… It will again.” closes alone; '
     '“given not by blood but by belonging” moved to the Living Line; hedging thinned '
     'to ~one per chapter; the American name-catalog cut to a gesture; the “could not '
     'have imagined” anaphora trimmed and the Faith virtue-gloss cut; body de-bolded; '
     'em-dashes down ~30%.')
write('Chapter Seven — your mom — is now three brief paragraphs, deliberately (she’s '
      'living). If you want one more sentence of public-safe biography (where she grew '
      'up, her work, anything she’d be glad to see), supply it and it goes in; '
      'otherwise the brevity reads as respect.')

# ---------------- McIver ----------------
chap('The McIvers — long-read (edits applied)')
done('The Ch 5 “scandal in two parishes” is now staked honestly up front instead of '
     'retracted in a late parenthetical; the research changelog moved to the Notes '
     '(the Vimy “wrong twice over” stays — there the correcting is the emotion); the '
     'twelve-child list no longer pre-spoils the empty chair; the ending recap cut by a '
     'third; the Notes’ broken emphasis markup repaired; superlatives halved.')
decide('Pre-existing quirk: the Ch 5 pull-quote (“A McIver met a Campbell. And held.”) '
       'echoes a line that lives at the end of Ch 4. Leave as a deliberate echo-ahead, '
       'or re-point the pull-quote at a Ch 5 line?')

# ---------------- Haiste ----------------
chap('The Haistes — long-read (edits applied)')
done('The prologue no longer spoils the whole book (Melanie, the marriage, and the '
     'convergence now arrive at the Living Line); Ch 3 retitled “the last quiet '
     'decades” (the old title promised “precision” the chapter never contained); '
     'doorkomen returned to the Lakemans — “They came through” now detonates exactly '
     'once, as the final line; three of the four “everything depends on this one man” '
     'climaxes cut (Ernest keeps his); the untethered survey-history trimmed; Ch 6 '
     'compressed to its verified baptism.')
write('Jim Haiste’s working life is the page’s acknowledged placeholder. What did '
      'Melanie’s grandfather actually do? Anything the family tells about him — a '
      'trade, a habit, a phrase — would anchor Ch 10 with something real.')
write('Records hunt (or a family ask): Ernest’s wife is unnamed in the manuscript, and '
      'she anchors the next two generations.')

# ---------------- Lakeman ----------------
chap('The Lakemans — long-read (edits applied)')
done('VERIFIED FIXES: the Suzanna/Rienk birthday gap is six days (25 Sep → 1 Oct 1918) '
     '— Ch 9 said five; Pieter Sijmonsz’s baptism is 23 Aug 1783 per the record audit '
     '— the prose said April. Also: the lake-house payoff now lands once, in Ch 9; '
     'Ch 7 retitled “the orphan who left the polder” (it duplicated Ch 9’s title); '
     'the polder-stillness thesis stated once; sibling lists moved to the table notes; '
     'Ch 10 now ends on Martin’s borderless childhood, not an administrative line.')
decide('One word to adjudicate: Ch 4 says all six children were baptized “in the '
       'Beemster Reformed Church,” but Pieter’s verified baptism is at Beets. Likely '
       '“Beemster and nearby parishes” — but it’s your record call.')
decide('Two pre-existing pull-quote quirks: Ch 6’s json line (“The whole line now hung '
       'on a seven-year-old boy”) doesn’t quite match the prose (“ran through that one '
       'orphaned boy”), and the Interlude’s (“A needle, a trade, and a century of '
       'life”) is a composed line, not an echo. Both render fine — align or leave?')
write('Martin’s chapter is still the thinnest chapter about the closest person. One or '
      'two memories of your father in your words — how he talked, what he drove, a day '
      'you remember — and Claude will set them in. This is the single highest-value '
      'paragraph you could write in the whole heritage set.')

# ---------------- Verboom ----------------
chap('The Verbooms — long-read (edits applied)')
done('The dek now leads with Suzanna and the needle; the Ch 3–4 name-rolls moved to '
     'the Notes (nothing deleted from the record); the trades now thread to the '
     'prologue’s “small enough to pack and useful enough to keep”; the Kinderdijk '
     'sentence split; the piece ends on the needle.')
decide('The deep-ancestor names (Tiggelman, Koppenaal, Hofland sides) now live in the '
       'Notes — happy there, or would you rather they appear as lineage-table rows?')
decide('Ch 1’s pull-quote (“He taught his daughter to sew; she became a seamstress for '
       'life.”) is a paraphrase of the prose, not an exact echo. Fine as is?')

# ---------------- Rycroft ----------------
chap('The Rycrofts — long-read (edits applied)')
done('The reach-diet: 25 → ~9, so the motif sings again; the itinerary told once; the '
     'Essex relitigation cut to the prologue + Notes; the chapter-wonder advertisements '
     'muted; Eric’s era-wallpaper cut; the Steinke insert cut to the verified frame; '
     'the de-bold sweep done (“The slip said Rycroft.” keeps its bold); the ending now '
     'lands on Faith Lana Lyn keeping her grandmother’s name.')
write('Eric (Ch 5) owns exactly two facts: born in Honolulu, and broke Alberta snow '
      'with his own boots. Does Melanie’s side remember ANYTHING else about him — a '
      'habit, a story at the family table, what he farmed? One detail gives the '
      'thinnest rung a third plank.')
write('Same ask for Henry and Martha Steinke (the Ch 6 insert): one remembered thing — '
      'occupation, church, the language at home — replaces what the cut textbook '
      'paragraph used to fake.')

# ---------------- Steinke ----------------
chap('The Steinkes — long-read (edits applied)')
done('The hold/durchhalten motif thinned to three beats; Ch 1 compressed to its '
     'documented floor; the chain recital cut from Ch 2; the mind-reading clause and '
     'the all-motifs catalogue cut; the kickers rewritten to carry dates and places '
     'instead of echoing their headings.')
decide('The prologue still defines durchhalten via “a Pomeranian or Volhynian-German '
       'farm family” — the only unverified-origin phrasing left in the body (the Notes '
       'say that link is NOT confirmed). Soften to “a German farm family”?')
write('This line’s gap is research, not writing: Bette’s 1959 Sexsmith marriage record '
      'and the 1931 census would unlock Henry and Martha’s origins in one stroke. Say '
      'the word when you want that chase run.')

# ---------------- closing ----------------
chap('How to send it back')
note('Any format works: scribble answers under each item here, or fire rough lists at '
     'Claude in chat. Everything gets unfolded into your voice, grounded to exactly what '
     'you give — expand the cadence, never invent the facts.')

out = r'C:\Users\thoma\Desktop\Heritage research\Your-Words-Needed.docx'
d.save(out)
print('saved:', out)
print('paragraphs:', len(d.paragraphs))
